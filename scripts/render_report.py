#!/usr/bin/env python3
"""Render V3.1 connected review prose with inline numeric citations and ACS-like references."""
from __future__ import annotations

import argparse
import json
import os
import sys
from collections import OrderedDict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from report_contract import TLDR_LABELS
from run_store import canonical, load_list, read, report_quality, validate, write

CN_FONT = os.environ.get("REPORT_CJK_FONT", "SimSun")
LATIN_FONT = "Times New Roman"


def set_font(style, size: float, bold: bool = False) -> None:
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), CN_FONT)


def styles_for(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(.9)
    section.top_margin = section.bottom_margin = Inches(.8)
    body = doc.styles["Normal"]
    set_font(body, 11)
    body.paragraph_format.line_spacing = 1.35
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.widow_control = True
    for name, size, before, after in [("Title", 19, 0, 14), ("Heading 1", 15, 16, 8),
                                     ("Heading 2", 12.5, 12, 6)]:
        style = doc.styles[name]
        set_font(style, size, name != "Title")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style in doc.styles:
        if style.type == 1:
            for border in list(style.element.xpath(".//w:pBdr")):
                border.getparent().remove(border)
    if "TLDR Label" not in doc.styles:
        from docx.enum.style import WD_STYLE_TYPE
        doc.styles.add_style("TLDR Label", WD_STYLE_TYPE.PARAGRAPH)
    label = doc.styles["TLDR Label"]
    label.base_style = doc.styles["Normal"]
    set_font(label, 11.5, True)
    label.paragraph_format.keep_with_next = True
    label.paragraph_format.space_after = Pt(6)


def citation_group(numbers: list[int]) -> str:
    numbers = sorted(set(numbers))
    groups, i = [], 0
    while i < len(numbers):
        end = i
        while end + 1 < len(numbers) and numbers[end + 1] == numbers[end] + 1:
            end += 1
        if end - i >= 2:
            groups.append(str(numbers[i]) + "–" + str(numbers[end]))
        else:
            groups.extend(str(n) for n in numbers[i:end + 1])
        i = end + 1
    return "[" + ",".join(groups) + "]" if groups else ""


def source_priority(source: dict) -> tuple:
    levels = {"FULLTEXT": 4, "PARTIAL": 3, "ABSTRACT": 2, "METADATA": 1}
    return (source.get("source_kind") != "secondary", levels.get(source.get("access_level"), 0),
            source.get("retrieval_surface", "HOST_OPEN_PAGE") == "HOST_OPEN_PAGE", source.get("accessed_date", ""))


def joined_citation(numbers: list[int]) -> str:
    # A zero-width Word Joiner keeps the citation attached to preceding punctuation.
    return "\u2060" + citation_group(numbers) if numbers else ""


class CitationIndex:
    def __init__(self, project: Path):
        self.project = project
        self.papers = {p["id"]: p for p in load_list(project, "papers")}
        self.sources = {s["source_id"]: s for s in load_list(project, "sources")}
        self.evidence = {e["evidence_id"]: e for e in load_list(project, "evidence")}
        self.claims = {c["claim_id"]: c for c in load_list(project, "claims")}
        self.order = OrderedDict()
        self.supporting_sources = {}

    def numbers(self, claim_ids: list[str]) -> list[int]:
        result = []
        for cid in claim_ids:
            for eid in self.claims[cid]["evidence_ids"]:
                evidence = self.evidence[eid]
                pid = canonical(self.project, evidence["paper_id"])
                source = self.sources[evidence["source_id"]]
                self.supporting_sources.setdefault(pid, {})[source["source_id"]] = source
                if pid not in self.order:
                    self.order[pid] = len(self.order) + 1
                if self.order[pid] not in result:
                    result.append(self.order[pid])
        return sorted(result)

    def entries(self) -> list[dict]:
        result = []
        for pid, number in self.order.items():
            sources = [s for s in self.supporting_sources[pid].values() if s.get("read_state") == "READ"]
            if not sources:
                raise ValueError("Cited work lacks a read supporting source: " + pid)
            paper = self.papers[pid]
            preferred = sorted(sources, key=source_priority, reverse=True)[0]
            web = paper.get("record_type") == "web"
            required = ("title",) if web else ("authors", "title", "venue", "year", "volume", "doi")
            missing = [key for key in required if not paper.get(key)]
            if not web and not (paper.get("pages") or paper.get("article_number")):
                missing.append("pages_or_article_number")
            result.append({"number": number, "paper_id": pid, "record_type": paper.get("record_type", "paper"),
                           "paper": paper, "preferred_source": preferred, "supporting_sources": sources,
                           "missing_bibliographic_fields": missing})
        return result


def add_citation(paragraph, numbers: list[int]) -> None:
    if numbers:
        run = paragraph.add_run(joined_citation(numbers))
        run.font.superscript = False
        run.font.name = LATIN_FONT


def paragraph(doc: Document, block: dict, citations: CitationIndex, label: str = "") -> None:
    para = doc.add_paragraph()
    if label:
        run = para.add_run(label + "：")
        run.bold = True
    if block.get("spans"):
        for span in block["spans"]:
            para.add_run(str(span.get("text", "")))
            add_citation(para, citations.numbers(span.get("claim_ids", [])))
    else:
        para.add_run(str(block.get("text", "")).strip())
        add_citation(para, citations.numbers(block.get("claim_ids", [])))


def table(doc: Document, block: dict, citations: CitationIndex, number: int) -> None:
    caption = doc.add_paragraph()
    caption.paragraph_format.keep_with_next = True
    caption.add_run("表 " + str(number) + "  " + block.get("caption", "研究路线比较")).bold = True
    cols = block["columns"]
    grid = doc.add_table(rows=1, cols=len(cols))
    grid.alignment, grid.autofit = WD_TABLE_ALIGNMENT.CENTER, False
    weights = block.get("column_weights") or [max(1, min(4, len(str(c)) / 3)) for c in cols]
    if len(weights) != len(cols) or any(float(x) <= 0 for x in weights):
        raise ValueError("Invalid table column weights")
    total = sum(float(x) for x in weights)
    widths = [Inches(6.7 * float(x) / total) for x in weights]
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:" + side)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    grid._tbl.tblPr.append(borders)
    for i, col in enumerate(cols):
        grid.columns[i].width = widths[i]
        cell = grid.rows[0].cells[i]
        cell.width, cell.text = widths[i], str(col)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "E7EFF5")
        cell._tc.get_or_add_tcPr().append(shade)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    grid.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for item in block["rows"]:
        cells = grid.add_row().cells
        for i, value in enumerate(item["cells"]):
            cells[i].width, cells[i].text = widths[i], str(value)
        add_citation(cells[-1].paragraphs[-1], citations.numbers(item.get("claim_ids", [])))
    short_table = len(grid.rows) <= 6 and sum(len(str(v)) for row in block["rows"] for v in row["cells"]) < 650
    for row_number, row in enumerate(grid.rows):
        if short_table:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                if short_table and row_number < len(grid.rows) - 1:
                    p.paragraph_format.keep_with_next = True
                for run in p.runs:
                    run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def bibliography(doc: Document, entries: list[dict]) -> None:
    doc.add_heading("参考文献", level=1)
    for entry in entries:
        paper, source = entry["paper"], entry["preferred_source"]
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Inches(.28)
        para.paragraph_format.first_line_indent = Inches(-.28)
        para.add_run("[" + str(entry["number"]) + "] ")
        authors = paper.get("authors_acs") or paper.get("authors") or paper.get("organization") or []
        author_text = "; ".join(str(a) for a in authors) if isinstance(authors, list) else str(authors)
        if author_text:
            para.add_run(author_text.rstrip(".") + ". ")
        para.add_run(paper["title"].rstrip(".") + ". ")
        if paper.get("record_type") == "web":
            para.add_run(source["url"] + " (accessed " + source["accessed_date"] + ").")
            continue
        venue = paper.get("journal_abbreviation") or paper.get("venue")
        if venue:
            para.add_run(str(venue)).italic = True
            para.add_run(" ")
        if paper.get("year"):
            para.add_run(str(paper["year"])).bold = True
        if paper.get("volume"):
            para.add_run(", ")
            para.add_run(str(paper["volume"])).italic = True
        page = paper.get("pages") or paper.get("article_number")
        if page:
            para.add_run(", " + str(page))
        if venue or paper.get("year"):
            para.add_run(". ")
        if paper.get("doi"):
            para.add_run("https://doi.org/" + paper["doi"] + ".")
        else:
            para.add_run(source["url"] + " (accessed " + source["accessed_date"] + ").")


def compose(project: Path, output: Path) -> dict:
    problems = validate(project, final=True)
    if problems:
        raise ValueError("Final evidence/narrative gate failed:\n- " + "\n- ".join(problems))
    report = read(project / "report_draft.json")
    citations = CitationIndex(project)
    template = Path(__file__).resolve().parent.parent / "assets" / "report-template.docx"
    doc = Document(template) if template.exists() else Document()
    styles_for(doc)
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(report["title"])
    doc.add_paragraph("太长不看版 / TL;DR", style="TLDR Label")
    for block, label in zip(report["tldr"], TLDR_LABELS):
        paragraph(doc, block, citations, label)
    table_count = 0
    for n, section in enumerate(report["sections"], 1):
        doc.add_heading(str(n) + ". " + section["title"], level=1)
        if section["kind"] in ("introduction", "conclusion"):
            for block in section["blocks"]:
                paragraph(doc, block, citations)
            continue
        paragraph(doc, section["opening"], citations)
        for j, sub in enumerate(section["subsections"], 1):
            doc.add_heading(f"{n}.{j} " + sub["title"], level=2)
            for block in sub["blocks"]:
                if block.get("type", "paragraph") == "table":
                    table_count += 1
                    table(doc, block, citations, table_count)
                else:
                    paragraph(doc, block, citations)
        paragraph(doc, section["closing"], citations)
    entries = citations.entries()
    bibliography(doc, entries)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    write(project / "report_citations.json", {"format_version": "3.1", "entries": entries})
    quality = report_quality(project)
    write(project / "report_quality.json", quality)
    return {"output": str(output), "references": len(entries), "summary_tables": table_count,
            "numbered_chapters": len(report["sections"]), "quality_warnings": quality["warnings"]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    try:
        print(json.dumps(compose(args.project.resolve(), args.output.resolve()), ensure_ascii=False, indent=2))
    except (ValueError, KeyError, OSError, TypeError) as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
