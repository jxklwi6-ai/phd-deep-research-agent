#!/usr/bin/env python3
"""Verify the V3.1 Word narrative, inline references, bibliography, and no-graphics contract."""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from report_contract import report_blocks, text_of
from render_report import CitationIndex, joined_citation, bibliography
from report_contract import TLDR_LABELS
from run_store import load_list, read, report_quality, validate

CITATION = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]")


def expanded(group: str) -> list[int]:
    result = []
    for part in group.split(","):
        bounds = re.split(r"[–-]", part.strip())
        if len(bounds) == 1:
            result.append(int(bounds[0]))
        else:
            start, end = map(int, bounds)
            if end < start or end - start > 10000:
                raise ValueError("Invalid numeric reference range")
            result.extend(range(start, end + 1))
    return result


def check(project: Path, docx: Path) -> dict:
    problems = validate(project, final=True)
    quality = report_quality(project)
    if problems:
        return {"passed": False, "problems": problems, "warnings": quality["warnings"]}
    if not docx.is_file() or docx.stat().st_size < 2000:
        return {"passed": False, "problems": problems + ["Missing or empty DOCX"]}
    document = Document(docx)
    with zipfile.ZipFile(docx) as archive:
        if any(name.startswith("word/media/") for name in archive.namelist()):
            problems.append("An image/media asset is embedded in the DOCX")
    if document.inline_shapes or document.element.xpath(".//w:drawing | .//w:pict"):
        problems.append("An image or drawing is present")
    report = read(project / "report_draft.json", {})
    headings = [p.text.strip() for p in document.paragraphs if p.style.name.startswith("Heading")]
    expected = []
    for n, section in enumerate(report.get("sections", []), 1):
        expected.append(str(n) + ". " + section["title"])
        for j, sub in enumerate(section.get("subsections", []), 1):
            expected.append(f"{n}.{j} " + sub["title"])
    expected.append("参考文献")
    if headings != expected:
        problems.append("DOCX headings do not match the mandatory V3.1 narrative outline")
    if any(p.style.name in ("Heading 3", "Heading 4") for p in document.paragraphs):
        problems.append("Fragmented H4 or deeper headings are not allowed")
    paragraphs = [p for p in document.paragraphs if p.text.strip()]
    if not paragraphs or paragraphs[0].style.name != "Title" or paragraphs[0].text != report.get("title"):
        problems.append("The specific title must appear once in the Word Title style")
    if sum(p.style.name == "Title" for p in document.paragraphs) != 1:
        problems.append("Exactly one report title is required")
    if len(paragraphs) < 2 or paragraphs[1].text != "太长不看版 / TL;DR" or paragraphs[1].style.name.startswith("Heading"):
        problems.append("Unnumbered TL;DR must immediately follow the title")
    if any(p.style.name == "Title" and p._p.xpath(".//w:pBdr") for p in document.paragraphs):
        problems.append("Decorative title borders are not allowed")
    if any(s.name == "Title" and s.element.xpath(".//w:pBdr") for s in document.styles):
        problems.append("A decorative line survives in the title style")

    bibliography_started, refs, body_citations, first_seen = False, [], [], []
    reference_paragraphs = {}
    all_text = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            para = Paragraph(element, document._body)
            all_text.append(para.text)
            if para.text.strip() == "参考文献":
                bibliography_started = True
                continue
            if bibliography_started and para.text.strip():
                match = re.match(r"^\[(\d+)\]\s+(.+)", para.text.strip())
                if not match:
                    problems.append("Malformed reference entry or stray paragraph after bibliography")
                else:
                    number = int(match.group(1))
                    refs.append(number)
                    reference_paragraphs[number] = para
                continue
            candidates = [para]
        elif element.tag == qn("w:tbl"):
            candidates = [p for row in Table(element, document._body).rows for cell in row.cells
                          for p in cell.paragraphs]
            all_text.extend(p.text for p in candidates)
            if bibliography_started:
                problems.append("A table appears after the bibliography heading")
        else:
            continue
        for para in candidates:
            for match in CITATION.finditer(para.text):
                numbers = expanded(match.group(1))
                body_citations.extend(numbers)
                for number in numbers:
                    if number not in first_seen:
                        first_seen.append(number)
            if any(r.font.superscript and CITATION.search(r.text) for r in para.runs):
                problems.append("References must be baseline bracketed numbers, not V3 superscripts")
    if refs != list(range(1, len(refs) + 1)) or first_seen != list(range(1, len(refs) + 1)):
        problems.append("References must be consecutive in first-citation order")
    if set(refs) != set(body_citations):
        problems.append("Body citation and bibliography numbers do not reconcile")
    # Reconstruct exact visible claim spans and table rows from the current evidence store.
    # Aggregate citation counts alone cannot detect a citation moved to another fact.
    index = CitationIndex(project)
    def expected_paragraph(block: dict, label: str = "") -> str:
        result = label + "：" if label else ""
        if block.get("spans"):
            for span in block["spans"]:
                result += str(span.get("text", "")) + joined_citation(index.numbers(span.get("claim_ids", [])))
        else:
            result += str(block.get("text", "")).strip() + joined_citation(index.numbers(block.get("claim_ids", [])))
        return result
    expected_body = [("paragraph", report["title"]), ("paragraph", "太长不看版 / TL;DR")]
    for block, label in zip(report["tldr"], TLDR_LABELS):
        expected_body.append(("paragraph", expected_paragraph(block, label)))
    table_no = 0
    for n, section in enumerate(report["sections"], 1):
        expected_body.append(("paragraph", str(n) + ". " + section["title"]))
        if section["kind"] in ("introduction", "conclusion"):
            expected_body.extend(("paragraph", expected_paragraph(block)) for block in section["blocks"])
            continue
        expected_body.append(("paragraph", expected_paragraph(section["opening"])))
        for j, sub in enumerate(section["subsections"], 1):
            expected_body.append(("paragraph", f"{n}.{j} " + sub["title"]))
            for block in sub["blocks"]:
                if block.get("type", "paragraph") != "table":
                    expected_body.append(("paragraph", expected_paragraph(block)))
                    continue
                table_no += 1
                expected_body.append(("paragraph", "表 " + str(table_no) + "  " + block.get("caption", "研究路线比较")))
                rows = [list(map(str, block["columns"]))]
                for row in block["rows"]:
                    cells = list(map(str, row["cells"]))
                    cells[-1] += joined_citation(index.numbers(row.get("claim_ids", [])))
                    rows.append(cells)
                expected_body.append(("table", rows))
        expected_body.append(("paragraph", expected_paragraph(section["closing"])))
    actual_body = []
    for element in document.element.body:
        if element.tag == qn("w:p"):
            text = Paragraph(element, document._body).text
            if text.strip() == "参考文献":
                break
            if text.strip():
                actual_body.append(("paragraph", text))
        elif element.tag == qn("w:tbl"):
            actual_body.append(("table", [[cell.text for cell in row.cells]
                                          for row in Table(element, document._body).rows]))
    if actual_body != expected_body:
        problems.append("DOCX prose/table content or sentence-level citations disagree with the current draft and evidence mapping")
    if len(document.tables) != quality["metrics"]["comparison_tables"]:
        problems.append("A required comparison table is missing or extra")
    whole_text = "\n".join(all_text)
    for block in report_blocks(report):
        if block.get("type") != "table":
            # Citations split text spans, so inspect each span separately.
            pieces = [str(s.get("text", "")) for s in block.get("spans", [])] or [text_of(block)]
            if any(piece.strip() and piece.strip() not in whole_text for piece in pieces):
                problems.append("A required synthesis paragraph/claim span is missing from the DOCX")
    mapping = read(project / "report_citations.json", {}).get("entries", [])
    if len(mapping) != len(refs):
        problems.append("Rendered reference count disagrees with citation-to-paper mapping")
    current_order = [(pid, number) for pid, number in index.order.items()]
    if [(entry["paper_id"], entry["number"]) for entry in mapping] != current_order:
        problems.append("Citation-to-paper mapping is stale relative to current claim evidence")
    current_entries = index.entries()
    if mapping != current_entries:
        problems.append("Citation source/provenance or bibliographic metadata is stale relative to the current evidence store")
    expected_bib = Document()
    bibliography(expected_bib, current_entries)
    expected_reference_texts = {entry["number"]: paragraph.text
                               for entry, paragraph in zip(current_entries, expected_bib.paragraphs[1:])}
    current_papers = {p["id"]: p for p in load_list(project, "papers")}
    metadata_gaps = []
    for entry in mapping:
        number = entry["number"]
        paper = current_papers.get(entry["paper_id"])
        para = reference_paragraphs.get(number)
        if not paper or not para:
            problems.append("A mapped bibliography work is missing")
            continue
        if para.text != expected_reference_texts.get(number):
            problems.append("Reference authors/pages/metadata disagree with the complete verified bibliographic record")
        if paper["title"] not in para.text:
            problems.append("Reference title disagrees with verified paper identity")
        if paper.get("doi") and ("https://doi.org/" + paper["doi"]) not in para.text:
            problems.append("Formal paper reference omits its verified DOI")
        if paper.get("record_type") == "web":
            if not re.search(r"\(accessed \d{4}-\d{2}-\d{2}\)", para.text):
                problems.append("Web reference requires an ISO access date")
        else:
            venue = paper.get("journal_abbreviation") or paper.get("venue")
            if venue and not any(r.text == str(venue) and r.italic for r in para.runs):
                problems.append("ACS-like journal name must be italic")
            if paper.get("year") and not any(r.text == str(paper["year"]) and r.bold for r in para.runs):
                problems.append("ACS-like publication year must be bold")
            if paper.get("volume") and not any(r.text == str(paper["volume"]) and r.italic for r in para.runs):
                problems.append("ACS-like volume must be italic")
        if entry.get("missing_bibliographic_fields"):
            metadata_gaps.append({"paper_id": entry["paper_id"], "missing": entry["missing_bibliographic_fields"]})
    if len(document.sections) != 1:
        problems.append("Unexpected multiple page-layout sections")
    else:
        page = document.sections[0]
        if abs(page.page_width / 914400 - 8.5) > .05 or abs(page.page_height / 914400 - 11) > .05:
            problems.append("Default report page size must remain Letter")
    return {"passed": not problems, "problems": sorted(set(problems)), "warnings": quality["warnings"],
            "reference_count": len(refs), "inline_citation_mentions": len(body_citations),
            "word_tables": len(document.tables), "narrative_metrics": quality["metrics"],
            "bibliographic_gaps": metadata_gaps, "has_images": bool(document.inline_shapes)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    args = parser.parse_args()
    try:
        result = check(args.project.resolve(), args.docx.resolve())
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if not result["passed"]:
            sys.exit(1)
    except (OSError, ValueError, KeyError, TypeError, zipfile.BadZipFile) as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
