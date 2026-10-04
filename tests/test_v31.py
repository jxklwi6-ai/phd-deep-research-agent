"""Synthetic audit fixtures, not scientific literature or a model writing example."""
from __future__ import annotations

import copy
import csv
import re
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document

SKILL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL / "scripts"))
import run_store as store
from package_skill import package
from quality_gate import check
from render_report import citation_group, compose
from report_contract import REVIEW_KEYS, TLDR_ROLES


def prose(text, refs=None, roles=None, content_kind="scientific"):
    out = {"type": "paragraph", "text": text, "claim_ids": refs or [], "content_kind": content_kind}
    if roles:
        out["roles"] = roles
    return out


class ResearchSkillEndToEnd(unittest.TestCase):
    def setUp(self):
        self.scratch = tempfile.TemporaryDirectory()
        self.addCleanup(self.scratch.cleanup)
        self.root = Path(self.scratch.name)
        self.project = self.root / "run"
        store.init(self.project, "Synthetic evidence-reading fixture", "2026-10-03", "zh-CN")
        self.make_corpus()

    def make_corpus(self):
        papers = [
            {"id": "P_OPEN", "title": "Synthetic open flow reaction", "authors": ["Tester, A.", "Researcher, B."],
             "year": 2025, "venue": "Test Journal", "volume": "12", "pages": "101-108",
             "doi": "10.1234/synthetic-open", "record_type": "paper", "role": "core",
             "screening_state": "include", "identity_checked": True, "first_seen_round": 1,
             "role_assigned_round": 1, "public_urls": ["https://example.org/open"]},
            {"id": "P_CLOSED", "title": "Synthetic closed catalyst result", "authors": ["Tester, C."],
             "year": 2026, "venue": "Example Journal", "volume": "7", "article_number": "e100",
             "doi": "10.1234/synthetic-closed", "record_type": "paper", "role": "core",
             "screening_state": "include", "identity_checked": True, "first_seen_round": 1,
             "role_assigned_round": 1, "public_urls": ["https://example.org/abstract"]}]
        store.upsert_papers(self.project, papers)
        store.upsert_generic(self.project, "sources", [
            {"source_id": "S_OPEN", "paper_id": "P_OPEN", "url": "https://example.org/open",
             "title": papers[0]["title"], "accessed_date": "2026-10-03", "access_level": "FULLTEXT",
             "read_state": "READ", "source_kind": "primary", "retrieval_surface": "HOST_OPEN_PAGE"},
            {"source_id": "S_CLOSED", "paper_id": "P_CLOSED", "url": "https://example.org/abstract",
             "title": papers[1]["title"], "accessed_date": "2026-10-03", "access_level": "ABSTRACT",
             "read_state": "READ", "source_kind": "publisher", "retrieval_surface": "HOST_OPEN_PAGE"}])
        store.upsert_generic(self.project, "evidence", [
            {"evidence_id": "E_OPEN", "paper_id": "P_OPEN", "source_id": "S_OPEN",
             "locator": "Results, paragraph 2", "evidence_level": "FULLTEXT",
             "paraphrase": "Synthetic text states that a flow condition was tested.", "direction": "support"},
            {"evidence_id": "E_CLOSED", "paper_id": "P_CLOSED", "source_id": "S_CLOSED",
             "locator": "Abstract, sentence 3", "evidence_level": "ABSTRACT",
             "paraphrase": "Synthetic abstract states that a catalyst comparison was made.", "direction": "support"}])
        store.upsert_generic(self.project, "claims", [
            {"claim_id": "C_OPEN", "claim_type": "summary", "standing": "supported",
             "text": "The synthetic open text tested a condition.", "evidence_ids": ["E_OPEN"]},
            {"claim_id": "C_CLOSED", "claim_type": "summary", "standing": "supported",
             "text": "The synthetic abstract states a comparison.", "evidence_ids": ["E_CLOSED"]}])
        for eid, family, route, query, candidates in [
                ("Q1", "broad", "keyword", "flow evidence", ["P_OPEN", "P_CLOSED"]),
                ("Q2", "recent", "recent", "recent catalyst evidence", ["P_CLOSED"])]:
            store.append_event(self.project, {"event_id": eid, "query_id": eid.lower(),
                "round": 1, "family": family, "route": route, "source": "host-web",
                "query": query, "status": "COMPLETED", "observed_returned_count": len(candidates),
                "candidate_ids": candidates})
        store.upsert_generic(self.project, "coverage_matrix", [
            {"subquestion": "main", "gap_severity": "important",
             "gap": "Synthetic fixture has no completed citation tracing",
             "direct_primary_evidence": "mixed", "next_query": "forward citation route"}])
        store.upsert_generic(self.project, "rounds", [{"round": 1, "query_ids": ["q1", "q2"],
            "new_concepts": ["synthetic fixture"], "new_contradictions": [], "new_terms": []}])
        limitation = "本用例使用合成记录，没有执行真实文献检索，不能据此判断领域覆盖或科学结论。"
        refs = ["C_OPEN", "C_CLOSED"]
        tldr_texts = [
            "本用例考察公开正文与闭源摘要两类合成记录，目标是检查不同阅读层级如何进入一篇连续的综述。",
            "合成正文支持已见条件的描述，合成摘要仅支持其中明确报告的比较；两种材料的可核查范围应保持区别。",
            "测试重点包括论文身份去重、逐句引文映射和可编辑比较表，这些是程序验证项目而非领域研究进展。",
            limitation,
            "正文先建立两类材料的整体关系，再比较证据能够支持的认识，最后说明证据边界如何影响综合判断。"]
        intro = [
            prose("证据阅读的共同任务是让读者理解判断所依赖的材料及其边界。这里的两份合成记录分别代表公开正文与公开摘要，服务于报告结构和引用实现的检查。", refs, ["object_value"]),
            prose("正文与摘要并非相同的阅读对象。正文可以提供已见段落的条件与比较，而摘要仅能支持实际显示的结论，不能补出其中没有给出的细节。", refs, ["principles_history"]),
            prose("本用例把两种记录纳入同一文献集合，以检查形式上分散的来源能否组织成围绕证据问题的连续论述，避免将每份记录单独写成摘要卡片。", refs, ["development_status"]),
            prose("当前需要解决的矛盾是，来源在主题上可以重要，但可见内容的完整程度不同；因此综合判断必须同时考虑信息的相关性和可验证范围。", refs, ["bottleneck"]),
            prose("检索截止日期为 2026-10-03。" + limitation + "正文先介绍两类记录的格局，再比较可支持的判断，最后讨论未读内容造成的限制。", [], ["scope_roadmap"], "method")]
        compare = {"type": "table", "caption": "两类合成来源的可见内容比较",
                   "columns": ["来源", "可见材料", "判断边界"], "column_weights": [1, 2, 3],
                   "rows": [{"cells": ["公开记录", "相关正文段落", "仅支持已读条件的描述"], "claim_ids": ["C_OPEN"]},
                            {"cells": ["闭源记录", "公开摘要", "不能补充未显示的实验细节"], "claim_ids": ["C_CLOSED"]}]}
        gap_table = {"type": "table", "caption": "来源边界与可核查问题",
                     "columns": ["问题", "已见依据", "尚不能确定"], "column_weights": [1, 2, 3],
                     "rows": [{"cells": ["正文记录", "合成条件段落", "该条件在其他体系的适用范围"], "claim_ids": ["C_OPEN"]},
                              {"cells": ["摘要记录", "合成比较结论", "比较所需的完整控制实验"], "claim_ids": ["C_CLOSED"]}]}
        def chapter(kind, title, opening, subs, closing):
            return {"kind": kind, "title": title, "opening": prose(opening, refs),
                    "subsections": subs, "closing": prose(closing, refs)}
        overview = chapter("overview", "两类阅读材料与总体格局",
            "两类记录围绕相近的证据问题组织，但提供的信息深度不同。先建立这一整体关系，才能判断后续的比较在哪里成立、在哪里需要限制。",
            [{"title": "公开正文提供的依据", "blocks": [prose("公开记录包含被实际阅读的条件段落，因此其描述可追溯至指定位置；这一性质使其适合支持已经显示的具体判断。", ["C_OPEN"]), compare]},
             {"title": "公开摘要的位置", "blocks": [prose("摘要记录给出比较的概括性结论，足以保留其在主题中的位置，却不能独立恢复未显示的细节。", ["C_CLOSED"])]}],
            "两种来源放在一起说明，主题相关性与阅读层级应分别记录。下一章据此比较支持范围，而不以全文可得性代替论文的重要性。")
        theme = chapter("theme", "比较判断与证据条件",
            "证据比较需要关注结论与条件之间的对应关系。公开正文和公开摘要可以共同进入综合判断，但应分别说明各自究竟支持了什么。",
            [{"title": "判断与原始位置的对应", "blocks": [{"type": "paragraph", "spans": [
                {"text": "正文记录支持已见条件的描述。", "claim_ids": ["C_OPEN"]},
                {"text": "摘要记录支持其中报告的比较。", "claim_ids": ["C_CLOSED"]}]}]},
             {"title": "不同阅读层级的比较", "blocks": [prose("两类记录能够比较的是可见结论，而不能把摘要中未展示的内容作为已经核查的事实；这也限定了综合段应采用的措辞。", refs)]}],
            "将判断与证据位置联系起来后，来源差异能够成为解释边界的一部分。仍未显示的控制和条件则需要留到局限部分，不能靠叙事补齐。")
        limits = chapter("limitations", "当前局限与可验证问题",
            "目前最主要的限制来自可见文本的范围，而非两类记录在主题上是否相关。综合时应保留摘要结论，同时明确无法独立检验的内容。",
            [{"title": "未显示内容的限制", "blocks": [prose("摘要中没有显示的实验条件和控制细节仍属未知，因此不能据此构建更强的机制结论或跨体系判断。", ["C_CLOSED"]), gap_table]},
             {"title": "后续核查的边界", "blocks": [prose(limitation, [], content_kind="method")]}],
            "这些限制将可支持的认识压缩到实际显示的内容，也提示后续核查必须针对缺少的条件和控制。报告的最后综合应据此收束，而不扩张结论。")
        conclusion = [
            prose("两类合成记录能够放入一条共同的论述主线，但其可核查范围并不相同。整体结构因此应先建立问题，再比较证据，最后把边界纳入结论。", refs, ["field_state"]),
            prose("最可靠的综合认识是，相关性决定材料在主题中的位置，阅读深度决定结论能够说到什么程度；二者需要共同约束报告的表达。", refs, ["scientific_synthesis"]),
            prose("尚未显示的条件与控制是当前判断的主要限制，不能把记录之间形式上的一致视为独立验证，也不能据此推断真实研究领域的状态。", refs, ["key_unresolved"]),
            prose("后续核查若能获得公开可见的相关条件和控制，将有助于判断概括性比较的适用范围。这一方向来自前述信息缺口，最终仍应以实际证据决定综合结论的强度。", refs, ["grounded_outlook"])]
        draft = {"format_version": "3.1", "title": "2025–2026 年合成证据记录阅读比较测试",
                 "as_of_date": "2026-10-03", "length_profile": "user_concise",
                 "length_reason": "A concise synthetic regression artifact rather than a literature report.",
                 "execution_mode": "SYNTHETIC_TEST_FIXTURE", "limitations": [limitation],
                 "tldr": [dict(prose(text, refs if role not in ("limits", "reader_takeaway") else [],
                                      content_kind="scientific" if role not in ("limits", "reader_takeaway") else "method"),
                               role=role) for role, text in zip(TLDR_ROLES, tldr_texts)],
                 "sections": [{"kind": "introduction", "title": "引言", "blocks": intro},
                              overview, theme, limits,
                              {"kind": "conclusion", "title": "结论与展望", "blocks": conclusion}]}
        store.write(self.project / "report_draft.json", draft)
        store.write(self.project / "narrative_plan.json", {
            "main_storyline": "Different visible texts establish different evidentiary boundaries.",
            "core_tension": "Relevance does not establish full-text verification.",
            "chapter_questions": ["What are the source types?", "What supports the comparison?", "What remains unseen?"],
            "overall_insights": ["Separate relevance and access", "Link facts to visible text", "Preserve unknown details"]})
        store.write(self.project / "editorial_review.json", {"format_version": "3.1", "items": {
            key: {"status": "PASS", "basis": "Synthetic fixture verifies the contract item " + key,
                  "locations": ["synthetic report"]} for key in REVIEW_KEYS}})

    def test_closed_core_and_duplicate_doi_keep_one_work(self):
        result = store.upsert_papers(self.project, [{"id": "P_DUP", "title": "Synthetic open flow reaction",
            "doi": "https://doi.org/10.1234/synthetic-open", "public_urls": ["https://example.org/mirror"]}])
        self.assertEqual(result["actions"][0]["canonical"], "P_OPEN")
        self.assertEqual(len(store.load_list(self.project, "papers")), 2)
        self.assertEqual(store.validate(self.project, final=True), [])

    def test_abstract_cannot_support_unread_detailed_claim(self):
        claims = store.load_list(self.project, "claims")
        claims[1]["claim_type"] = "detailed"
        store.write(self.project / "claims.json", claims)
        self.assertTrue(any("lacks read public full text" in p for p in store.validate(self.project, True)))

    def test_missing_introduction_conclusion_or_closing_is_rejected(self):
        original = store.read(self.project / "report_draft.json")
        for index in (0, -1):
            draft = copy.deepcopy(original)
            draft["sections"].pop(index)
            store.write(self.project / "report_draft.json", draft)
            self.assertTrue(store.validate(self.project, True))
        draft = copy.deepcopy(original)
        del draft["sections"][2]["closing"]
        store.write(self.project / "report_draft.json", draft)
        self.assertTrue(any("closing synthesis" in p for p in store.validate(self.project, True)))

    def test_old_v3_draft_requires_rewriting(self):
        store.write(self.project / "report_draft.json", {"title": "old", "summary": "old", "sections": []})
        self.assertTrue(any("rewrite a V3 draft" in p for p in store.validate(self.project, True)))

    def test_default_full_length_requires_substance_or_recorded_exception(self):
        draft = store.read(self.project / "report_draft.json")
        draft["length_profile"] = "full"
        store.write(self.project / "report_draft.json", draft)
        self.assertTrue(any("6000-12000" in p for p in store.validate(self.project, True)))

    def test_resume_cannot_silently_ignore_a_changed_cutoff(self):
        before = (self.project / "manifest.json").read_bytes()
        with self.assertRaisesRegex(ValueError, "different search cutoff"):
            store.init(self.project, "Synthetic evidence-reading fixture", "2026-10-04", "zh-CN")
        self.assertEqual((self.project / "manifest.json").read_bytes(), before)

    def test_docx_acs_format_sentence_citations_and_audit_package(self):
        output = self.project / "report.docx"
        result = compose(self.project, output)
        self.assertEqual(result["references"], 2)
        self.assertTrue(check(self.project, output)["passed"])
        doc = Document(output)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("正文记录支持已见条件的描述。[1]摘要记录支持其中报告的比较。[2]", text.replace("\u2060", ""))
        self.assertIn("Tester, A.; Researcher, B.", text)
        self.assertIn("https://doi.org/10.1234/synthetic-open", text)
        self.assertNotIn("Works cited", text)
        self.assertNotIn("检索与证据审计", text)
        self.assertNotIn("实际检索事件", text)
        store.export(self.project)
        with zipfile.ZipFile(self.project / "research_audit.zip") as z:
            for item in ("retrieval_ledger.jsonl", "narrative_plan.json", "editorial_review.json",
                         "report_quality.json", "report_citations.json", "research_audit.txt"):
                self.assertIn(item, z.namelist())
            self.assertIn("实际检索事件：2", z.read("research_audit.txt").decode("utf-8"))
        with (self.project / "literature_map.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual({r["id"]: r["highest_read_level"] for r in rows},
                         {"P_OPEN": "FULLTEXT", "P_CLOSED": "ABSTRACT"})
        zipped = self.root / "skill.zip"
        package(SKILL, zipped)
        with zipfile.ZipFile(zipped) as z:
            self.assertIn("phd-deep-research-agent-v3-1/LICENSE", z.namelist())
            self.assertIn("phd-deep-research-agent-v3-1/references/report-output-spec-v3.1.md", z.namelist())
            self.assertFalse(any("__pycache__" in n for n in z.namelist()))

    def test_indexed_abstract_never_becomes_fulltext(self):
        sources = store.load_list(self.project, "sources")
        sources[1].update({"retrieval_surface": "HOST_SEARCH_RESULT", "tool_result_ref": "test-search-ref",
                           "abstract_complete": True})
        store.write(self.project / "sources.json", sources)
        self.assertEqual(store.validate(self.project, True), [])
        output = self.project / "report.docx"
        compose(self.project, output)
        mapping = store.read(self.project / "report_citations.json")
        self.assertEqual(mapping["entries"][1]["preferred_source"]["retrieval_surface"], "HOST_SEARCH_RESULT")
        sources[1]["access_level"] = "FULLTEXT"
        store.write(self.project / "sources.json", sources)
        self.assertTrue(any("cannot establish public full text" in p for p in store.validate(self.project, True)))

    def test_source_mismatch_and_secondary_only_science_fail(self):
        sources = store.load_list(self.project, "sources")
        sources[1]["source_kind"] = "secondary"
        store.write(self.project / "sources.json", sources)
        self.assertTrue(any("only secondary webpage" in p for p in store.validate(self.project, True)))

    def test_unknown_sentence_claim_and_invisible_limit_are_rejected(self):
        draft = store.read(self.project / "report_draft.json")
        draft["sections"][2]["subsections"][0]["blocks"][0]["spans"][0]["claim_ids"] = ["C_UNKNOWN"]
        draft["limitations"].append("A limitation hidden only in metadata.")
        store.write(self.project / "report_draft.json", draft)
        errors = store.validate(self.project, True)
        self.assertTrue(any("unknown claim ID" in p for p in errors))
        self.assertTrue(any("exists only in metadata" in p for p in errors))

    def test_incomplete_editorial_review_is_rejected(self):
        review = store.read(self.project / "editorial_review.json")
        review["items"]["continuous_readability"]["status"] = "PENDING"
        store.write(self.project / "editorial_review.json", review)
        self.assertTrue(any("continuous_readability" in p for p in store.validate(self.project, True)))

    def test_removed_table_or_moved_citations_fail_the_docx_gate(self):
        output = self.project / "report.docx"
        compose(self.project, output)
        altered = Document(output)
        for table in list(altered.tables):
            table._tbl.getparent().remove(table._tbl)
        damaged = self.project / "missing-tables.docx"
        altered.save(damaged)
        self.assertFalse(check(self.project, damaged)["passed"])
        altered = Document(output)
        for para in altered.paragraphs:
            if para.text == "参考文献":
                break
            for run in para.runs:
                run.text = re.sub(r"\[\d+(?:[,–-]\d+)*\]", "", run.text)
        for table in altered.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.text = re.sub(r"\[\d+(?:[,–-]\d+)*\]", "", run.text)
        altered.paragraphs[2].add_run("[1,2]")
        moved = self.project / "moved-citations.docx"
        altered.save(moved)
        self.assertFalse(check(self.project, moved)["passed"])

    def test_changed_claim_mapping_invalidates_an_old_docx(self):
        output = self.project / "report.docx"
        compose(self.project, output)
        draft = store.read(self.project / "report_draft.json")
        def replace_refs(value):
            if isinstance(value, dict):
                for key in value:
                    if key == "claim_ids" and value[key]:
                        value[key] = ["C_OPEN"]
                    else:
                        replace_refs(value[key])
            elif isinstance(value, list):
                for item in value:
                    replace_refs(item)
        replace_refs(draft)
        store.write(self.project / "report_draft.json", draft)
        self.assertFalse(check(self.project, output)["passed"])

    def test_reference_author_or_page_tampering_fails(self):
        output = self.project / "report.docx"
        compose(self.project, output)
        altered = Document(output)
        for para in altered.paragraphs:
            if para.text.startswith("[1] "):
                for run in para.runs:
                    run.text = run.text.replace("Tester, A.; Researcher, B.", "Invented, Z.")
                    run.text = run.text.replace("101-108", "999-1000")
        altered.save(output)
        self.assertFalse(check(self.project, output)["passed"])
        with self.assertRaisesRegex(ValueError, "DOCX integrity failed"):
            store.export(self.project)

    def test_changed_source_provenance_invalidates_mapping_and_export(self):
        output = self.project / "report.docx"
        compose(self.project, output)
        sources = store.load_list(self.project, "sources")
        sources[0]["source_id"] = "S_OPEN_NEW"
        sources[0]["url"] = "https://example.org/new-public-version"
        store.write(self.project / "sources.json", sources)
        evidence = store.load_list(self.project, "evidence")
        evidence[0]["source_id"] = "S_OPEN_NEW"
        store.write(self.project / "evidence.json", evidence)
        self.assertEqual(store.validate(self.project, True), [])
        self.assertFalse(check(self.project, output)["passed"])
        with self.assertRaisesRegex(ValueError, "changed after rendering"):
            store.export(self.project)

    def test_web_reference_and_missing_metadata_do_not_get_invented(self):
        store.upsert_papers(self.project, [{"id": "P_CLOSED", "title": "Synthetic closed catalyst result",
                                          "record_type": "web", "doi": "", "organization": "Test Organization"}])
        # Existing DOI is deliberately retained by conservative upsert; remove it explicitly for this synthetic web fixture.
        papers = store.load_list(self.project, "papers")
        papers[1]["doi"] = ""
        papers[1]["authors"] = []
        store.write(self.project / "papers.json", papers)
        output = self.project / "web.docx"
        compose(self.project, output)
        self.assertTrue(check(self.project, output)["passed"])
        self.assertIn("(accessed 2026-10-03)", "\n".join(p.text for p in Document(output).paragraphs))

    def test_repeated_queries_do_not_fake_saturation(self):
        store.append_event(self.project, {"event_id": "REPEAT", "query_id": "q3", "round": 2,
            "family": "broad", "route": "keyword", "source": "host-web", "query": "flow evidence",
            "status": "COMPLETED", "observed_returned_count": 2, "candidate_ids": ["P_OPEN", "P_CLOSED"]})
        store.upsert_generic(self.project, "rounds", [{"round": 2, "query_ids": ["q3"],
            "new_concepts": [], "new_contradictions": [], "new_terms": []}])
        result = store.evaluate(self.project)["round_trajectory"][-1]
        self.assertEqual(result["fresh_query_route_source_combinations"], 0)
        self.assertFalse(result["eligible_low_novelty"])

    def test_two_new_low_novelty_rounds_require_route_and_anchor_coverage(self):
        families = store.read(self.project / "query_families.json")
        for item in families["families"]:
            if item["name"] not in ("broad", "recent", "citation", "synonym"):
                item.update({"status": "NOT_APPLICABLE", "reason": "Synthetic stopping-rule fixture"})
        store.write(self.project / "query_families.json", families)
        for n, family, route in [(2, "citation", "backward"), (3, "synonym", "forward"), (3, "synonym", "related")]:
            store.append_event(self.project, {"event_id": f"e{n}{route}", "query_id": f"q{n}{route}", "round": n,
                "family": family, "route": route, "source": "host-web", "query": f"synthetic {route} {n}",
                "status": "COMPLETED", "observed_returned_count": 1, "candidate_ids": ["P_OPEN"]})
        store.upsert_generic(self.project, "rounds", [{"round": n, "query_ids": [],
            "new_concepts": [], "new_contradictions": [], "new_terms": []} for n in (2, 3)])
        store.upsert_generic(self.project, "coverage_matrix", [{"subquestion": "main", "gap_severity": "none"}])
        store.upsert_generic(self.project, "anchor_audit", [
            {"audit_id": route, "subquestion": "main", "route": route, "status": "COMPLETED",
             "candidate_ids": ["P_OPEN"]} for route in store.ANCHOR_ROUTES])
        self.assertTrue(store.evaluate(self.project)["saturated"])

    def test_citation_groups(self):
        self.assertEqual(citation_group([7, 6]), "[6,7]")
        self.assertEqual(citation_group([5, 3, 4]), "[3–5]")
        self.assertEqual(citation_group([1, 2, 3, 6, 7]), "[1–3,6,7]")


if __name__ == "__main__":
    unittest.main()
